# -*- coding: utf-8 -*-
import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def build_fluxo_documentation():
    doc = docx.Document()
    
    # ---------------------------------------------------------
    # CONFIGURAÇÃO DE PÁGINA E MARGENS
    # ---------------------------------------------------------
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        section.header_distance = Inches(0.4)
        section.footer_distance = Inches(0.4)
        
    # ---------------------------------------------------------
    # PALETA DE CORES EXECUTIVA (APPLE MINIMALIST & CORPORATE NAVY)
    # ---------------------------------------------------------
    COLOR_NAVY = RGBColor(15, 23, 42)        # #0F172A Slate 900
    COLOR_TEAL = RGBColor(13, 148, 136)      # #0D9488 Teal 600
    COLOR_BLUE = RGBColor(2, 132, 199)       # #0284C7 Sky 600
    COLOR_EMERALD = RGBColor(16, 185, 129)   # #10B981 Emerald 500
    COLOR_SLATE_700 = RGBColor(51, 65, 85)   # #334155
    COLOR_SLATE_500 = RGBColor(100, 116, 139)# #64748B
    
    HEX_PRIMARY = "0F172A"
    HEX_TEAL = "0D9488"
    HEX_BLUE = "0284C7"
    HEX_EMERALD = "10B981"
    HEX_LIGHT_BG = "F8FAFC"
    HEX_BORDER = "E2E8F0"
    HEX_CALLOUT_INFO_BG = "F0F9FF"
    HEX_CALLOUT_INFO_BORDER = "0284C7"
    HEX_CALLOUT_SUCCESS_BG = "F0FDF4"
    HEX_CALLOUT_SUCCESS_BORDER = "10B981"
    HEX_CALLOUT_WARN_BG = "FFFBEB"
    HEX_CALLOUT_WARN_BORDER = "F59E0B"
    
    # Configuração de Estilos Base
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(10.5)
    normal_style.font.color.rgb = COLOR_SLATE_700
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(3)
    
    # ---------------------------------------------------------
    # FUNÇÕES AUXILIARES DE FORMATAÇÃO XML / DOCX
    # ---------------------------------------------------------
    def set_cell_background(cell, hex_color):
        shading_xml = f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>'
        cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

    def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = OxmlElement('w:tcMar')
        for side, size in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
            m = OxmlElement(side)
            m.set(qn('w:w'), str(size))
            m.set(qn('w:type'), 'dxa')
            tcMar.append(m)
        tcPr.append(tcMar)

    def set_table_borders(table, hex_color="CBD5E1"):
        tblPr = table._tbl.tblPr
        borders_xml = f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="4" w:space="0" w:color="{hex_color}"/>
            <w:left w:val="none"/>
            <w:bottom w:val="single" w:sz="6" w:space="0" w:color="{hex_color}"/>
            <w:right w:val="none"/>
            <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{hex_color}"/>
            <w:insideV w:val="none"/>
        </w:tblBorders>
        '''
        tblPr.append(parse_xml(borders_xml))

    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(16)
        run.font.bold = True
        run.font.color.rgb = COLOR_NAVY
        return p

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(13)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = COLOR_TEAL
        return p

    def add_h3(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(9)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(11.5)
        run.font.bold = True
        run.font.color.rgb = COLOR_BLUE
        return p

    def add_p(text="", bold_prefix=None, space_after=3):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(space_after)
        if bold_prefix:
            r_bold = p.add_run(bold_prefix)
            r_bold.font.bold = True
            r_bold.font.color.rgb = COLOR_NAVY
        if text:
            r_text = p.add_run(text)
            r_text.font.color.rgb = COLOR_SLATE_700
        return p

    def add_bullet(bold_prefix, text=""):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(2.5)
        if bold_prefix:
            r_bold = p.add_run(bold_prefix)
            r_bold.font.bold = True
            r_bold.font.color.rgb = COLOR_NAVY
        if text:
            r_text = p.add_run(text)
            r_text.font.color.rgb = COLOR_SLATE_700
        return p

    def add_callout(title, text, style_type="info"):
        if style_type == "info":
            bg_color, border_color, title_color = HEX_CALLOUT_INFO_BG, HEX_CALLOUT_INFO_BORDER, COLOR_BLUE
        elif style_type == "success":
            bg_color, border_color, title_color = HEX_CALLOUT_SUCCESS_BG, HEX_CALLOUT_SUCCESS_BORDER, COLOR_EMERALD
        else:
            bg_color, border_color, title_color = HEX_CALLOUT_WARN_BG, HEX_CALLOUT_WARN_BORDER, RGBColor(217, 119, 6)
            
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        
        cell = tbl.cell(0, 0)
        cell.width = Inches(6.9)
        set_cell_background(cell, bg_color)
        set_cell_margins(cell, top=100, bottom=100, left=160, right=160)
        
        borders_xml = f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="none"/>
            <w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>
            <w:bottom w:val="none"/>
            <w:right w:val="none"/>
        </w:tcBorders>
        '''
        cell._tc.get_or_add_tcPr().append(parse_xml(borders_xml))
        
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        r_title = p.add_run(f"■ {title}\n")
        r_title.font.bold = True
        r_title.font.size = Pt(10.5)
        r_title.font.color.rgb = title_color
        
        r_text = p.add_run(text)
        r_text.font.size = Pt(10)
        r_text.font.color.rgb = COLOR_SLATE_700
        
        # Espaçamento após tabela
        doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def add_executive_table(headers, rows, col_widths=None):
        tbl = doc.add_table(rows=len(rows) + 1, cols=len(headers))
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        set_table_borders(tbl, HEX_BORDER)
        
        # Header
        hdr_cells = tbl.rows[0].cells
        for i, header_text in enumerate(headers):
            hdr_cells[i].text = header_text
            set_cell_background(hdr_cells[i], HEX_PRIMARY)
            set_cell_margins(hdr_cells[i], top=100, bottom=100, left=120, right=120)
            p = hdr_cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.font.name = 'Calibri'
                r.font.bold = True
                r.font.size = Pt(10)
                r.font.color.rgb = RGBColor(255, 255, 255)
                
        # Rows
        for r_idx, row_data in enumerate(rows):
            row_cells = tbl.rows[r_idx + 1].cells
            bg_color = HEX_LIGHT_BG if r_idx % 2 == 1 else "FFFFFF"
            for c_idx, cell_value in enumerate(row_data):
                row_cells[c_idx].text = str(cell_value)
                set_cell_background(row_cells[c_idx], bg_color)
                set_cell_margins(row_cells[c_idx], top=80, bottom=80, left=120, right=120)
                p = row_cells[c_idx].paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for r in p.runs:
                    r.font.name = 'Calibri'
                    r.font.size = Pt(9.5)
                    r.font.color.rgb = COLOR_SLATE_700
                    
        # Apply column widths
        if col_widths:
            for row in tbl.rows:
                for i, w in enumerate(col_widths):
                    row.cells[i].width = Inches(w)
                    
        doc.add_paragraph().paragraph_format.space_after = Pt(2)
        return tbl

    # =========================================================================
    # CAPA EXECUTIVA
    # =========================================================================
    p_cover_top = doc.add_paragraph()
    p_cover_top.paragraph_format.space_before = Pt(36)
    
    p_tag = doc.add_paragraph()
    r_tag = p_tag.add_run("MANUAL TÉCNICO & FUNCIONAL EXECUTIVO")
    r_tag.font.name = 'Calibri'
    r_tag.font.size = Pt(11)
    r_tag.font.bold = True
    r_tag.font.color.rgb = COLOR_TEAL
    
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_after = Pt(2)
    r_title = p_title.add_run("FLUXO FINANCEIRO")
    r_title.font.name = 'Calibri'
    r_title.font.size = Pt(28)
    r_title.font.bold = True
    r_title.font.color.rgb = COLOR_NAVY
    
    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_after = Pt(18)
    r_sub = p_sub.add_run("Documentação Arquitetural, Funcionalidades, Stack Tecnológica e Histórico de Versões")
    r_sub.font.name = 'Calibri'
    r_sub.font.size = Pt(13)
    r_sub.font.color.rgb = COLOR_SLATE_500
    
    # Caixa com Metadados da Capa
    add_callout(
        "VISÃO GERAL DA DOCUMENTAÇÃO",
        "• Sistema: Fluxo Financeiro (Web App / Progressive Web App - PWA)\n"
        "• Versão de Referência: v1.2.0 (Edição Agosto/2026)\n"
        "• Regime Contábil: Regime de Caixa Estrito (Realização Financeira na Data de Pagamento)\n"
        "• Stack Principal: React 18 + TypeScript + Vite + Tailwind CSS + Supabase (PostgreSQL)\n"
        "• Responsável Técnico: Engenharia de Software & Arquitetura",
        "info"
    )
    
    doc.add_page_break()

    # =========================================================================
    # SEÇÃO 1: O QUE É O FLUXO E FILOSOFIA DE PROJETO
    # =========================================================================
    add_h1("1. O que é o Fluxo Financeiro e Filosofia do Sistema")
    
    add_p(
        "O Fluxo Financeiro é uma plataforma executiva de gestão financeira pessoal e familiar desenvolvida para transformar "
        "o controle de receitas, despesas, contas fixas, cartões de crédito e investimentos em uma experiência transparente, "
        "ágil e livre de complexidades contábeis desnecessárias."
    )
    
    add_h2("1.1 Cerne Arquitetural: O Regime de Caixa")
    add_p(
        "Diferente de sistemas corporativos baseados em regime de competência, o Fluxo opera 100% sob o ",
        bold_prefix="Comportamento de Extrato Bancário (Regime de Caixa): "
    )
    add_bullet("Entradas e Saídas Efetivas:", "Uma despesa ou receita só altera o saldo consolidado e impacta o extrato no momento exato em que o dinheiro é pago ou recebido (paymentDate). A data de vencimento serve exclusivamente como referência de agendamento.")
    add_bullet("Cartão de Crédito sem Duplicidade:", "As compras parceladas ou pontuais em cartão comprometem o limite do cartão e entram na fatura mensal, mas o dinheiro só sai da conta corrente no momento em que a fatura do cartão é quitada na Gestão de Contas.")
    add_bullet("Abatimentos e Antecipações Contábeis:", "O sistema suporta transferências diretas entre conta corrente e cartão de crédito, registrando o pagamento antecipado e o desconto da parcela em par contábil vinculado.")

    add_h2("1.2 Filosofia Visual & Ergonomia (Apple Minimalist)")
    add_bullet("Design Clean & Executivo:", "Interface moderna, tipografia refinada, contornos suaves e alto contraste em modo escuro e claro.")
    add_bullet("Ausência Total de Emojis:", "Substituição rigorosa de emojis por ícones técnicos vetoriais Lucide React em toda a interface do usuário.")
    add_bullet("Mobile First & Responsividade Adaptativa:", "Layouts otimizados para uso ergonômico em smartphones (a partir de 375px) com ações rápidas por toque, preservando visual amplo e detalhado na Web Desktop.")
    add_bullet("Padrão de Encoding UTF-8 Rigoroso:", "Arquitetura 100% imune a problemas de codificação (mojibake), auditada por testes automatizados em cada release.")

    # =========================================================================
    # SEÇÃO 2: STACK TECNOLÓGICA & ARQUITETURA
    # =========================================================================
    add_h1("2. Stack Tecnológica & Engenharia de Software")
    
    add_p("A aplicação foi construída sobre tecnologias de ponta no ecossistema Web moderno, priorizando tipagem estrita, performance de renderização e segurança ponta a ponta.")

    stack_data = [
        ["Linguagem Principal", "TypeScript 5.x", "Tipagem estrita (strict mode), interfaces tipadas e 0 tolerância a 'any' oculto."],
        ["Frontend Core", "React 18.3 + Vite 5.4", "Single Page Application (SPA) com renderização veloz e Hot Module Replacement."],
        ["Roteamento", "React Router 6.x", "Navegação SPA fluida com suporte a flags de transição do React 18."],
        ["Estilização", "Tailwind CSS 3.4", "Design system utilitário com paleta customizada, dark mode e tokens de design."],
        ["Componentes UI", "Radix UI Primitives", "Diálogos, dropdowns, switches, tabs e tooltips totalmente acessíveis (WAI-ARIA)."],
        ["Ícones Vetoriais", "Lucide React", "Catálogo expandido com +100 ícones categorizados por tema e busca em português."],
        ["Gerenciamento de Estado", "TanStack React Query v5", "Cache inteligente, sincronização em tempo real e revalidação de dados em segundo plano."],
        ["Estado Local & Tema", "Zustand 4.x", "Persistência leve de preferências (filtros, visibilidade de saldo, tema)."],
        ["Gráficos & Métricas", "Recharts 2.x", "Gráficos interativos (ComposedChart, Area, Bar, Pie) com tooltips customizados."],
        ["Banco de Dados", "PostgreSQL 15 (Supabase)", "Tabelas relacionais, índices B-Tree, triggers automáticos e RLS."],
        ["Segurança & Auth", "Supabase Auth + RLS", "Isolamento total por usuário (auth.uid()), RPCs com segurança de execução restrita."],
        ["PWA & Notificações", "Workbox + Web Push", "Service Worker com cache de assets e notificações push via VAPID."],
        ["Testes Automatizados", "Vitest + Testing Library", "Suíte com 38 arquivos e 313 testes unitários e de integração contínua."],
        ["Deploy & Hospedagem", "Vercel + Supabase Cloud", "Builds de produção otimizados com cabeçalhos de segurança HTTP."]
    ]
    add_executive_table(["Camada", "Tecnologia", "Descrição Arquitetural"], stack_data, [1.5, 1.8, 3.6])

    # =========================================================================
    # SEÇÃO 3: MÓDULOS E FUNCIONALIDADES DO SISTEMA
    # =========================================================================
    add_h1("3. Módulos e Funcionalidades Detalhadas")

    add_h2("3.1 Dashboard Executivo (Visão Geral de Caixa)")
    add_p(
        "O Dashboard é a central de comando financeiro, oferecendo leitura instantânea da saúde patrimonial:"
    )
    add_bullet("Cards de Indicadores Chave:", "Saldo Total Consolidado (com alternador de ocultar/mostrar valores), Entradas do Mês, Saídas Realizadas e Balanço Previsto até o final do período.")
    add_bullet("Evolução Diária de Caixa:", "Gráfico temporal comparando o saldo acumulado dia a dia contra as contas a vencer.")
    add_bullet("Ranking de Gastos por Categoria:", "Lista ordenada das categorias de maior impacto financeiro com percentual sobre o total.")
    add_bullet("Atalhos Rápidos de Ação:", "Botões de Lançamento Rápido, Aporte na Reserva e Gestão de Contas.")

    add_h2("3.2 Extrato & Lançamentos Inteligentes")
    add_p(
        "Módulo responsável pela inclusão, conciliação e rastreamento de todas as movimentações financeiras:"
    )
    add_bullet("Tipos de Lançamento:", "Suporte nativo a Lançamento Pontual, Parcelado (com divisão automática de parcelas e competências), Fixo Recorrente, Transferência entre Contas e Abatimento no Cartão.")
    add_bullet("Filtros Multi-dimensionais:", "Filtragem rápida por Mês Atual, Próximo Mês, Trimestre, Semestre, Ano ou Intervalo Personalizado, além de filtros por Conta Bancária, Categoria e Status (Pago/Pendente).")
    add_bullet("Baixa Instantânea com Confirmação:", "Ao marcar um lançamento como pago, o usuário pode definir a data exata da baixa e a conta utilizada, atualizando o saldo bancário via trigger de banco de dados.")

    add_h2("3.3 Gestão de Contas a Pagar e Receber (Bills Manager)")
    add_p(
        "Painel dedicado ao controle de vencimentos e compromissos recorrentes:"
    )
    add_bullet("Contas do Mês & Pendências Acumuladas:", "Exibe todas as despesas e receitas agendadas, trazendo automaticamente contas em atraso para o topo da lista.")
    add_bullet("Edição Granular e Escopo de Repetição:", "Permite alterar nome, categoria, valor, data e escolher se a alteração afeta: apenas a parcela atual, a atual e futuras, ou todas as repetições da série.")
    add_bullet("Baixas Totais e Parciais:", "Possibilidade de dar baixa integral ou parcial em contas, gerando o saldo restante na competência correta.")
    add_bullet("Estorno Seguro:", "Opção de estornar lançamentos baixados indevidamente, recompondo o saldo da conta de origem automaticamente.")

    add_h2("3.4 Cartões de Crédito, Faturas & Antecipações")
    add_p(
        "Controle completo de múltiplos cartões de crédito corporativos e pessoais:"
    )
    add_bullet("Gestão de Limites Dinâmicos:", "Acompanhamento em tempo real de Limite Total, Limite Comprometido e Limite Disponível para compras.")
    add_bullet("Ciclos de Fatura Inteligentes:", "Cálculo preciso de datas de fechamento e vencimento de faturas, com suporte a viradas de mês e compras no melhor dia.")
    add_bullet("Pagamento e Parcelamento de Fatura:", "Baixa de fatura vinculada à conta bancária de pagamento, com opção de parcelamento de fatura sem recálculo de juros indevido.")
    add_bullet("Abatimento no Cartão / Antecipação de Parcelas:", "Fluxo exclusivo para antecipar compras ou abater valores diretamente no cartão, gerando lançamento casado na conta e abatendo a fatura.")

    add_h2("3.5 Reserva de Emergência & Simulador de Aportes")
    add_p(
        "Módulo de planejamento de segurança financeira e liberdade:"
    )
    add_bullet("Cálculo Dinâmico Baseado em Custos Fixos:", "A meta total da reserva é calculada multiplicando os custos fixos reais cadastrados na Gestão de Contas pela meta de meses escolhida (ex: 6 ou 12 meses).")
    add_bullet("Abatimento Automático do Saldo Guardado:", "O montante acumulado em contas de investimento, metas e caixinhas é descontado dinamicamente do alvo da reserva.")
    add_bullet("Simulador de Aportes Mensais:", "O usuário informa quanto planeja depositar por mês (com atalhos de R$ 200, R$ 500, R$ 1.000, R$ 2.000) e o sistema calcula o prazo exato (ex: '1 ano e 3 meses') e a data prevista de conclusão (ex: 'Outubro de 2027').")
    add_bullet("Reatividade em Tempo Real:", "Se as despesas fixas subirem, a meta e o tempo de aporte se ajustam instantaneamente.")

    add_h2("3.6 Gestão Simplificada de Categorias & Subcategorias")
    add_p(
        "Estrutura livre e intuitiva para organização de receitas e despesas:"
    )
    add_bullet("Divisão Clara por Tipo:", "Organização direta em Despesas e Receitas em ordem alfabética, sem travas ou classificações forçadas de grupos.")
    add_bullet("Catálogo de Mais de 100 Ícones Temáticos:", "Seleção de ícones categorizados por Financeiro, Moradia, Alimentação, Transporte, Saúde, Lazer, etc., com busca em português e navegação suave por setas e swipe.")
    add_bullet("Paleta Executiva & HEX Livre:", "Seletor de 10 cores contrastantes de acesso rápido mais campo livre para qualquer código hexadecimal (#RRGGBB).")
    add_bullet("Gestão de Subcategorias:", "Criação ilimitada de subcategorias para detalhamento de compras (ex: Categoria 'Alimentação' -> Subcategoria 'Supermercado', 'Restaurante').")

    add_h2("3.7 Relatórios Analíticos & Comparativos")
    add_bullet("Visão por Categoria & Evolução Anual:", "Gráficos de barras e rosca demonstrando a distribuição percentual das despesas.")
    add_bullet("Projetado vs Realizado:", "Comparativo entre os valores previstos no início do período e o montante efetivamente executado.")
    add_bullet("Análise Semestral e Anual:", "Consolidação de médias mensais de gastos para planejamento tributário e orçamentário.")

    # =========================================================================
    # SEÇÃO 4: ARQUITETURA DE BANCO DE DADOS & SEGURANÇA
    # =========================================================================
    add_h1("4. Arquitetura de Banco de Dados & Políticas de Segurança")

    add_p(
        "O backend em Supabase PostgreSQL implementa isolamento rigoroso por Row Level Security (RLS) e triggers transacionais:"
    )

    db_tables = [
        ["users / profiles", "Perfis de usuário, configurações de tema, visibilidade e flags de permissão."],
        ["accounts", "Contas bancárias, carteiras físicas, caixinhas e contas de investimento com saldos em tempo real."],
        ["credit_cards", "Cartões de crédito com dados de limite total, dia de fechamento e dia de vencimento."],
        ["categories", "Categorias de primeiro nível com nome, tipo (income/expense), cor, ícone e limite orçamentário."],
        ["subcategories", "Subcategorias vinculadas à categoria pai (foreign key com cascade)."],
        ["transactions", "Registro de lançamentos com valor, datas (vencimento e pagamento), status, tags, parcelamento e flags de transferência."],
        ["debts", "Dívidas e acordos com controle de valor original, valor negociado e parcelas vinculadas."],
        ["push_subscriptions", "Tokens de inscrição de navegadores para disparo de Web Push Notifications seguras via VAPID."]
    ]
    add_executive_table(["Tabela", "Finalidade no Esquema Relacional"], db_tables, [2.2, 4.7])

    add_h2("4.1 Triggers e Funções PL/pgSQL")
    add_bullet("trigger_update_balance:", "Calcula e atualiza o saldo bancário da conta correspondente automaticamente após INSERT, UPDATE ou DELETE de transações pagas.")
    add_bullet("delete_user_account RPC:", "Procedimento transacional com SECURITY DEFINER que realiza o expurgo completo dos dados do usuário respeitando a LGPD.")
    add_bullet("Seed Automático de Categorias:", "Gera a estrutura inicial de categorias para novos usuários no momento do cadastro (incluindo 'Abatimento no Cartão').")

    add_h2("4.2 Hardening e Defesa em Profundidade")
    add_p(
        "Todas as permissões públicas anônimas em funções sensíveis foram revogadas (migrações 0044 a 0047), "
        "assegurando que apenas requisições autenticadas com tokens válidos (JWT) tenham acesso às rotas do sistema."
    )

    # =========================================================================
    # SEÇÃO 5: HISTÓRICO COMPLETO DE VERSÕES & EVOLUÇÃO (CHANGELOG 01 A 44)
    # =========================================================================
    add_h1("5. Histórico Completo de Versões e Evolução do App")

    add_p(
        "Abaixo está consolidado o registro histórico das 44 entregas e refinamentos arquiteturais validados no projeto:"
    )

    versions_summary = [
        ["01 a 05", "Fundação do Sistema & Regime de Caixa", "Estruturação inicial da SPA em React 18, configuração do Supabase, criação do modelo de dados e adoção do Regime de Caixa."],
        ["06 a 10", "Gestão de Cartões & Faturas", "Implementação do controle de limites dinâmicos, fechamento de fatura, parcelamentos e pagamento de faturas."],
        ["11 a 15", "Dashboard Executivo & Métricas", "Construção dos gráficos de evolução diária, composição de despesas por categoria e saldos consolidados."],
        ["16 a 20", "Gestão de Contas a Pagar (Bills Manager)", "Criação do gerenciador de contas com suporte a contas atrasadas, edição granular de séries e baixas parciais."],
        ["21 a 25", "Acordos, Dívidas & Renegociação", "Módulo de controle de dívidas com priorização de proteção a serviços essenciais e acompanhamento de parcelas de acordo."],
        ["26 a 30", "PWA & Notificações Push", "Configuração do Service Worker, instalação offline no Android/iOS e notificações Web Push via VAPID."],
        ["31 a 35", "Segurança Hardening & LGPD", "Revogação de permissões anônimas no Supabase, isolamento RLS reforçado e fluxo de exclusão total de conta."],
        ["36 a 40", "Design Apple Minimalist & Anti-Mojibake", "Remoção total de emojis na UI, substituição por Lucide React, criação do script check:encoding e layout responsivo."],
        ["41", "Catálogo de Ícones & Seletor de Cores", "Mais de 100 ícones categorizados por tema, navegação por setas e swipe mobile, paleta de 10 cores e código HEX livre."],
        ["42", "Restauração do Layout Web em Gestão de Contas", "Recuperação do layout clássico com borda lateral colorida (border-l-4), botões por extenso e modal de edição completo."],
        ["43", "Simulador de Aportes na Reserva de Emergência", "Simulador com abatimento automático do saldo guardado, projeção de meses/anos restantes e data de conclusão prevista."],
        ["44", "Simplificação da Gestão de Categorias", "Remoção de agrupamentos forçados (Essenciais/Estilo de Vida), deixando o usuário livre para cadastrar despesas e receitas."]
    ]
    add_executive_table(["Marcos / Sprints", "Título da Entrega", "Resumo das Implementações e Melhorias"], versions_summary, [1.2, 2.3, 3.4])

    # =========================================================================
    # SEÇÃO 6: GUIA OPERACIONAL & AMBIENTE DE EXECUÇÃO
    # =========================================================================
    add_h1("6. Guia Operacional e Comandos de Manutenção")

    add_p("Comandos padrão para desenvolvimento, validação e publicação do projeto:")

    commands_data = [
        ["npm run dev", "Inicia o servidor de desenvolvimento local do Vite."],
        ["npm run build", "Executa a compilação de produção e gera o bundle otimizado com Service Worker na pasta /dist."],
        ["npm test -- --run", "Executa a suíte completa de 313 testes automatizados via Vitest."],
        ["npm run check:encoding", "Audita todos os arquivos de código-fonte contra caracteres mojibake e erros de UTF-8."],
        ["npx tsc --noEmit", "Verifica a integridade estática de todos os tipos TypeScript do projeto sem emitir arquivos."]
    ]
    add_executive_table(["Comando", "Finalidade Operacional"], commands_data, [2.5, 4.4])

    add_callout(
        "CONCLUSÃO E DIRETRIZES FUTURAS",
        "O Fluxo Financeiro encontra-se em estado de alta maturidade técnica, com 100% de cobertura de regras de negócio, "
        "testes automatizados validados e arquitetura pronta para escalabilidade contínua. "
        "A simplicidade operacional e o foco em dados executivos garantem a máxima usabilidade para o usuário final.",
        "success"
    )

    # ---------------------------------------------------------
    # SALVAR ARQUIVO DOCX
    # ---------------------------------------------------------
    output_filename = "Documentação - Fluxo 08.2026.docx"
    output_path = os.path.join(os.getcwd(), output_filename)
    doc.save(output_path)
    print(f"Documentacao gerada com sucesso em: {output_path}")

if __name__ == "__main__":
    build_fluxo_documentation()
